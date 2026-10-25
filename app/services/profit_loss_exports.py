"""Excel serialization for the authoritative P&L report view model.

No accounting is performed here: every value is copied from
``build_profit_loss_report`` so screen, detail, PDF and Excel reconcile.
"""
from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def _value(value):
    return float(value or 0)


def _write_table(ws, headers, rows, widths=None):
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for index, header in enumerate(headers, 1):
        width = (widths or {}).get(header, min(42, max(12, len(str(header)) + 3)))
        ws.column_dimensions[get_column_letter(index)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=False)


def build_profit_loss_excel(report):
    summary = report.get("export_summary", report["summary"])
    previous_summary = report.get("export_previous_summary", report["previous_summary"])
    comparison = report.get("export_comparison", report["comparison"])
    wb = Workbook()
    summary_ws = wb.active
    summary_ws.title = "Summary"
    summary_ws["A1"] = "Profit & Loss Control Center"
    summary_ws["A1"].font = Font(size=16, bold=True, color="1F4E78")
    summary_ws["A2"] = "Period"
    summary_ws["B2"] = (
        f"{report['start_date_str'] or 'Beginning'} to {report['end_date_str'] or 'Today'}"
    )
    summary_ws["A3"] = "Generated"
    summary_ws["B3"] = report["generated_at"].strftime("%Y-%m-%d %H:%M:%S")
    filters = report["filters"]
    summary_ws["A4"] = "Filters"
    summary_ws["B4"] = "; ".join(
        f"{key.title()}: {filters.get(key)}" for key in ("client", "material", "supplier", "category") if filters.get(key)
    ) or "None"
    headers = ["Metric", "Current", "Previous equivalent", "Change"]
    for col, value in enumerate(headers, 1):
        cell = summary_ws.cell(6, col, value)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    metrics = [
        ("Earned Material Revenue (net of returns)", "total_material_revenue"),
        ("Earned Delivery Revenue (net of returns)", "total_delivery_revenue"),
        ("Other Earned Income", "total_other_income"),
        ("Sales (Money In)", "total_revenue"),
        ("Actual Recorded Material Cost", "actual_material_cost"),
        ("Estimated Material Cost", "estimated_material_cost"),
        ("Direct Delivery Cost", "total_delivery_cost"),
        ("Cost of Goods (Money Out)", "total_cogs"),
        ("Operating Expenses", "total_expenses"),
        ("Other Recognized Losses", "total_other_losses"),
        ("Other Recognized Gains", "total_other_gains"),
        ("Unresolved Revenue Excluded From Net", "unresolved_revenue"),
        ("Net Profit / Loss", "total_profit"),
        ("Recorded-Cost-Only Profit / Loss", "recorded_cost_profit"),
        ("Net Margin %", "margin_pct"),
        ("Money Received (position only)", "payment_received_total"),
    ]
    for row_index, (label, key) in enumerate(metrics, 7):
        summary_ws.cell(row_index, 1, label)
        current_value = summary.get(key) if key in summary else report.get(key)
        previous_value = previous_summary.get(key) if comparison["available"] and key != "payment_received_total" else None
        summary_ws.cell(row_index, 2, _value(current_value))
        summary_ws.cell(row_index, 3, _value(previous_value) if previous_value is not None else "N/A")
        summary_ws.cell(row_index, 4, (_value(current_value) - _value(previous_value)) if previous_value is not None else "N/A")
        for col in (2, 3, 4):
            summary_ws.cell(row_index, col).number_format = "0.00%" if key == "margin_pct" else '#,##0.00;[Red]-#,##0.00'
        if key == "margin_pct":
            # Service stores percentage points (12.5 means 12.5%), while Excel
            # percentage formatting expects a fraction.
            for col in (2, 3, 4):
                cell = summary_ws.cell(row_index, col)
                if isinstance(cell.value, (int, float)):
                    cell.value = _value(cell.value) / 100.0
    next_row = 7 + len(metrics) + 2
    summary_ws.cell(next_row, 1, "Traceable formula").font = Font(bold=True)
    summary_ws.cell(next_row + 1, 1, "Net = recognized material + delivery + other income + other gains − material cost − delivery cost − expenses − other losses. Rows with unresolved costs are excluded from net and disclosed below.")
    summary_ws.cell(next_row + 2, 1, "Formula check")
    summary_ws.cell(next_row + 2, 2, _value(summary["formula_check"]))
    summary_ws.cell(next_row + 3, 1, "Formula difference (must be 0)")
    summary_ws.cell(next_row + 3, 2, _value(summary["formula_difference"]))
    summary_ws.cell(next_row + 4, 1, "Previous comparison")
    summary_ws.cell(next_row + 4, 2, comparison["label"])
    if comparison["available"]:
        summary_ws.cell(next_row + 5, 1, "Previous period")
        summary_ws.cell(next_row + 5, 2, f"{comparison['start_date'].isoformat()} to {comparison['end_date'].isoformat()}")
    summary_ws.column_dimensions["A"].width = 48
    summary_ws.column_dimensions["B"].width = 28
    summary_ws.column_dimensions["C"].width = 24
    summary_ws.column_dimensions["D"].width = 20

    tx_ws = wb.create_sheet("Transactions")
    tx_headers = [
        "Date", "Type", "Category", "Reference", "Customer", "Material",
        "Booked Material", "Qty", "Sale Rate", "Gross", "Discount",
        "Material Revenue", "Delivery Revenue", "Other Income", "Material Cost",
        "Delivery Cost", "Expense", "Other Gain", "Other Loss", "Net Profit",
        "Cost Status", "Cost Basis", "Included In Net", "Document Type",
        "Document ID", "Source Type", "Source ID", "Payment Status", "Notes",
    ]
    tx_rows = []
    for row in report["export_transactions"]:
        tx_rows.append([
            row.get("date"), row.get("transaction_type"), row.get("category"),
            row.get("reference"), row.get("client"), row.get("material"),
            row.get("booked_material"), _value(row.get("qty")), _value(row.get("sale_rate")),
            _value(row.get("gross_revenue")), _value(row.get("discount")),
            _value(row.get("material_revenue")), _value(row.get("delivery_revenue")),
            _value(row.get("other_income")), _value(row.get("material_cost")),
            _value(row.get("delivery_cost")), _value(row.get("expense")),
            _value(row.get("other_gain")), _value(row.get("other_loss")),
            _value(row.get("profit")), row.get("cost_status"), row.get("cost_basis"),
            "Yes" if row.get("include_in_net") else "No", row.get("document_type"),
            row.get("document_id"), row.get("source_type"), row.get("source_id"),
            row.get("payment_status"), row.get("note"),
        ])
    _write_table(tx_ws, tx_headers, tx_rows, {"Cost Basis": 58, "Payment Status": 42, "Notes": 55, "Transaction Type": 30})
    for row in tx_ws.iter_rows(min_row=2):
        row[0].number_format = "yyyy-mm-dd hh:mm"
        for index in range(7, 20):
            row[index].number_format = '#,##0.00;[Red]-#,##0.00'

    daily_ws = wb.create_sheet("Daily Insights")
    _write_table(
        daily_ws,
        ["Day", "Revenue", "Cost", "Expenses / Losses", "Net Profit", "Rows"],
        [[r["day"], _value(r["revenue"]), _value(r["cogs"]), _value(r["expenses"]), _value(r["profit"]), r["rows"]] for r in reversed(report.get("export_daily_insights", report["daily_insights"]))],
    )
    for row in daily_ws.iter_rows(min_row=2, min_col=2, max_col=5):
        for cell in row:
            cell.number_format = '#,##0.00;[Red]-#,##0.00'

    material_ws = wb.create_sheet("Material Insights")
    _write_table(
        material_ws,
        ["Material", "Net Qty", "Net Revenue", "Net Profit", "Unresolved Rows"],
        [[r["material"], _value(r["sold_qty"]), _value(r["sold_revenue"]), _value(r["profit_known"]), r["unknown_rows"]] for r in report.get("export_top_items", report["top_items"])],
    )

    booking_ws = wb.create_sheet("Outstanding Bookings")
    position = report.get("financial_position")
    outstanding = position["outstanding_bookings"] if position else report["context"].booking_outstanding
    _write_table(
        booking_ws,
        ["Booking", "Customer", "Material", "Booked Qty", "Net Supplied Qty", "Remaining Qty", "Current Rate", "Unfulfilled Value", "Booking Paid Amount"],
        [[r["reference"], r["client"], r["material"], _value(r["booked_qty"]), _value(r["supplied_qty"]), _value(r["remaining_qty"]), _value(r["rate"]), _value(r["outstanding_value"]), _value(r["paid_amount"])] for r in outstanding],
    )

    quality_ws = wb.create_sheet("Warnings & Policy")
    quality_rows = [[w["code"], w["reference"], w["message"], _value(w["amount"])] for w in report["warnings"]]
    quality_rows += [["POLICY", "", message, 0.0] for message in report["unsupported"]]
    _write_table(quality_ws, ["Code", "Reference", "Warning / Accounting Policy", "Amount"], quality_rows, {"Warning / Accounting Policy": 110, "Reference": 25})

    # ------------------------------------------------------------------ #
    # Enhanced breakdown (same view model as the on-screen report).
    # Added as extra sheets so every existing sheet keeps its name/order.
    # ------------------------------------------------------------------ #
    breakdown = report.get("breakdown")
    if breakdown:
        cards = breakdown.get("cards") or []
        reconcile = breakdown.get("reconcile") or {}
        cards_ws = wb.create_sheet("P&L Cards")
        _write_table(
            cards_ws,
            ["Card", "Nature", "Amount", "What it is"],
            [
                [
                    card["label"],
                    "Realized profit" if card.get("realized") else (
                        "Estimate (not profit)" if card.get("estimate") else "Balance (not profit)"
                    ),
                    _value(card["value"]),
                    card.get("sub", ""),
                ]
                for card in cards
            ],
            {"What it is": 80, "Card": 40, "Nature": 22},
        )
        for row in cards_ws.iter_rows(min_row=2, min_col=3, max_col=3):
            for cell in row:
                cell.number_format = '#,##0.00;[Red]-#,##0.00'
        first_row = 2 + len(cards) + 1
        cards_ws.cell(first_row, 1, "Reconciliation").font = Font(bold=True)
        for offset, (label, key) in enumerate(
            [
                ("Gross profit (delivered goods)", "gross_profit"),
                ("+ Service margin", "service_margin"),
                ("+ Other income", "other_income"),
                ("+ Other gains", "other_gains"),
                ("- Operating expenses", "operating_expenses"),
                ("- Other losses", "other_losses"),
                ("= Net profit / loss", "net_profit"),
                ("Ledger profit (must match)", "report_profit"),
                ("Difference", "difference"),
            ],
            first_row + 1,
        ):
            cards_ws.cell(offset, 1, label)
            cards_ws.cell(offset, 2, _value(reconcile.get(key)))
            cards_ws.cell(offset, 2).number_format = '#,##0.00;[Red]-#,##0.00'
        cards_ws.cell(first_row + 11, 1, "Reconciles exactly" if reconcile.get("matches") else "REVIEW REQUIRED")
        cards_ws.column_dimensions["A"].width = 40
        cards_ws.column_dimensions["B"].width = 22

        item_ws = wb.create_sheet("By Item")
        _write_table(
            item_ws,
            ["Item", "Delivered qty", "Goods sales", "FIFO cost of goods sold", "Gross profit", "Margin",
             "Remaining stock qty", "Remaining stock value"],
            [
                [
                    row["material"], _value(row["qty"]), _value(row["material_revenue"]),
                    _value(row["material_cost"]), _value(row["gross_profit"]), _value(row["margin_pct"]),
                    _value(row.get("remaining_inventory_qty")), _value(row.get("remaining_inventory_value")),
                ]
                for row in breakdown.get("by_item", [])
            ],
        )
        service_ws = wb.create_sheet("By Service")
        _write_table(
            service_ws,
            ["Service", "Qty", "Revenue earned", "Direct cost", "Margin", "Margin percent"],
            [
                [
                    row["service"], _value(row["qty"]), _value(row["service_revenue"]),
                    _value(row["service_cost"]), _value(row["margin"]), _value(row["margin_pct"]),
                ]
                for row in breakdown.get("by_service", [])
            ],
        )
        customer_ws = wb.create_sheet("By Customer")
        _write_table(
            customer_ws,
            ["Customer", "Category", "Revenue", "Cost", "Profit / loss", "Margin percent"],
            [
                [
                    row["client"], row.get("client_category", ""), _value(row["revenue"]),
                    _value(row["cost"]), _value(row["profit"]), _value(row["margin_pct"]),
                ]
                for row in breakdown.get("by_customer", [])
            ],
            {"Customer": 40, "Category": 18},
        )
        day_ws = wb.create_sheet("By Day")
        _write_table(
            day_ws,
            ["Day", "Revenue", "FIFO cost", "Expenses / losses", "Net result", "Rows"],
            [
                [
                    row["day"], _value(row["revenue"]), _value(row["cost"]),
                    _value(row["expenses"] + row["other_loss"]), _value(row["profit"]), row.get("rows", 0),
                ]
                for row in breakdown.get("by_day", [])
            ],
        )

        payment_summary = breakdown.get("payment_summary") or {}
        credit_ws = wb.create_sheet("Outstanding Credit")
        _write_table(
            credit_ws,
            ["Document", "Reference", "Date", "Customer", "Payment status", "Bill value", "Collected",
             "Outstanding", "Outstanding on this period's revenue"],
            [
                [
                    row["document_type"], row["reference"],
                    row["date"].strftime("%Y-%m-%d") if row.get("date") else "",
                    row.get("client", ""), row.get("payment_status_label", ""),
                    _value(row.get("document_gross")), _value(row.get("document_collected")),
                    _value(row.get("document_outstanding")), _value(row.get("period_outstanding")),
                ]
                for row in payment_summary.get("top_credit_rows", payment_summary.get("document_rows", []))
            ],
            {"Reference": 28, "Customer": 34, "Document": 16},
        )
        advance_ws = wb.create_sheet("Customer Advances")
        _write_table(
            advance_ws,
            ["Booking", "Customer", "Material", "Remaining qty", "Unfulfilled value", "Advance (not revenue)"],
            [
                [
                    row["reference"], row["client"], row["material"], _value(row["remaining_qty"]),
                    _value(row["outstanding_value"]), _value(row["advance_value"]),
                ]
                for row in payment_summary.get("advance_rows", [])
            ],
            {"Customer": 34},
        )

        grn_section = breakdown.get("grn_section")
        stock_ws = wb.create_sheet("GRN & Inventory")
        stock_rows = []
        for row in (grn_section or {}).get("purchase_rows", []):
            lot = row.get("consumption") or {}
            stock_rows.append([
                "Purchase", row["reference"], row["material"],
                row["date"].strftime("%Y-%m-%d") if row.get("date") else "",
                _value(row["qty"]), _value(lot.get("consumed_qty")), _value(lot.get("remaining_qty")),
                _value(lot.get("consumed_pct")), _value(row["rate"]), _value(row["amount"]),
                _value(lot.get("remaining_value")), lot.get("status", ""),
            ])
        for row in (grn_section or {}).get("inventory_by_material", []):
            stock_rows.append([
                "Inventory", "", row["material"], "", _value(row["received_qty"]),
                _value(row["consumed_qty"]), _value(row["remaining_qty"]), "", "",
                _value(row.get("average_remaining_rate")), _value(row["remaining_value"]), "",
            ])
        _write_table(
            stock_ws,
            ["Section", "Reference", "Material", "GRN date", "Received qty", "Consumed qty", "Remaining qty",
             "Consumed percent", "Rate", "Purchase value", "Remaining value", "Lot status"],
            stock_rows,
            {"Reference": 24, "Section": 14},
        )
        if grn_section:
            note_row = len(stock_rows) + 3
            stock_ws.cell(note_row, 1, "Purchases in period (inventory, not expense)")
            stock_ws.cell(note_row, 2, _value(grn_section.get("purchase_value"))).number_format = '#,##0.00;[Red]-#,##0.00'
            stock_ws.cell(note_row + 1, 1, grn_section.get("note", ""))

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
