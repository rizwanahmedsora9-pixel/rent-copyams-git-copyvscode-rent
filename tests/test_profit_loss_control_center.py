"""Fourteen isolated accounting regressions for the P&L Control Center.

Every test builds a throw-away SQLite database through the shared app factory.
No production data, deployment path, or live application database is touched.
"""
from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from werkzeug.security import generate_password_hash

from app.services.profit_loss import (
    build_profit_loss_report,
    build_transaction_detail,
    resolve_report_period,
)
from models import (
    Account,
    AccountTransaction,
    Booking,
    BookingAllocation,
    BookingItem,
    CashFlowEntry,
    Client,
    DeliveryPerson,
    DeliveryPersonPayment,
    DirectSale,
    DirectSaleItem,
    Entry,
    FbmCashDrawerEntry,
    GRN,
    GRNAllocation,
    GRNItem,
    Material,
    MaterialReturn,
    MaterialReturnItem,
    Payment,
    Service,
    Supplier,
    SupplierPayment,
    User,
    WaiveOff,
    db,
)


@pytest.fixture()
def pl_app(app_factory, tmp_path):
    app = app_factory(tmp_path / "profit_control_center.db")
    app.config["TESTING"] = True
    with app.app_context():
        db.create_all()
        from app.services.schema import _ensure_default_admin
        _ensure_default_admin()
        db.session.commit()
    return app


def login(client, username="Admin", password="Admin@fbm12345"):
    response = client.post("/login", data={"username": username, "password": password})
    assert response.status_code in (302, 303), response.get_data(as_text=True)[:400]


def seed_base(*, material_rate=100.0):
    material = Material(name="CEMENT", code="MAT-1", unit="Bags", unit_price=material_rate, total=10000)
    client = Client(name="CUSTOMER ONE", code="CL-1", category="General", is_active=True)
    supplier = Supplier(name="SUPPLIER ONE", is_active=True)
    account = Account(name="MAIN CASH", account_type="Cash", category="Cash", type="Cash", is_active=True)
    db.session.add_all([material, client, supplier, account])
    db.session.flush()
    return material, client, supplier, account


def add_grn(supplier, material, qty, rate, posted, bill="GRN-1"):
    grn = GRN(supplier_id=supplier.id, supplier=supplier.name, manual_bill_no=bill, date_posted=posted, is_void=False)
    db.session.add(grn)
    db.session.flush()
    item = GRNItem(grn_id=grn.id, mat_name=material.name, qty=qty, price_at_time=rate, is_void=False)
    db.session.add(item)
    db.session.flush()
    return grn, item


def add_sale(client, posted, items, *, category="Credit Customer", paid=0, discount=0, bill=None, delivery_cost=0, void=False):
    sale = DirectSale(
        client_name=client.name, client_code=client.code, category=category,
        amount=sum(qty * rate for _name, qty, rate, _cost in items), paid_amount=paid,
        discount=discount, manual_bill_no=bill, date_posted=posted,
        delivery_rent_cost=delivery_cost, is_void=void,
    )
    db.session.add(sale)
    db.session.flush()
    rows = []
    for name, qty, rate, cost in items:
        item = DirectSaleItem(
            sale_id=sale.id, product_name=name, qty=qty, price_at_time=rate,
            cost_rate_at_sale=cost,
        )
        db.session.add(item)
        db.session.flush()
        rows.append(item)
    return sale, rows


def add_booking(client, posted, items, *, paid=0, discount=0, bill="BK-1"):
    booking = Booking(
        client_name=client.name, amount=sum(qty * rate for _name, qty, rate in items),
        paid_amount=paid, discount=discount, manual_bill_no=bill,
        date_posted=posted, is_void=False,
    )
    db.session.add(booking)
    db.session.flush()
    rows = []
    for name, qty, rate in items:
        item = BookingItem(booking_id=booking.id, material_name=name, qty=qty, price_at_time=rate)
        db.session.add(item)
        db.session.flush()
        rows.append(item)
    return booking, rows


def report(start="2026-01-01", end="2026-12-31", **kwargs):
    return SimpleNamespace(**build_profit_loss_report(start_date=start, end_date=end, **kwargs))


def test_01_cash_sale_is_earned_before_and_independently_of_receipt(pl_app):
    with pl_app.app_context():
        material, client, _supplier, _account = seed_base(material_rate=100)
        sale, _items = add_sale(client, datetime(2026, 1, 10), [("CEMENT", 10, 150, None)], category="Cash", paid=1500)
        db.session.add(Payment(client_id=client.id, client_name=client.name, amount=500, payment_type="Receipt", date_posted=datetime(2026, 1, 20), is_void=False))
        db.session.commit()

        january = report("2026-01-01", "2026-01-31")
        assert january.summary["total_material_revenue"] == pytest.approx(1500)
        assert january.summary["total_cogs"] == pytest.approx(1000)
        assert january.summary["total_profit"] == pytest.approx(500)
        assert january.payment_received_total == pytest.approx(2000)
        assert not any(row["source"] == "Payment" for row in january.rows)

        receipt_day = report("2026-01-20", "2026-01-20")
        assert receipt_day.summary["total_profit"] == 0
        assert receipt_day.summary["total_revenue"] == 0
        assert receipt_day.payment_received_total == pytest.approx(500)
        assert sale.id


def test_02_credit_sale_purchase_and_supplier_payment_are_not_new_profit(pl_app):
    with pl_app.app_context():
        material, client, supplier, _account = seed_base(material_rate=999)
        _grn, grn_item = add_grn(supplier, material, 100, 80, datetime(2026, 1, 1))
        sale, items = add_sale(client, datetime(2026, 1, 5), [("CEMENT", 10, 120, 80)])
        items[0].grn_item_id = grn_item.id
        db.session.add(GRNAllocation(sale_id=sale.id, sale_item_id=items[0].id, grn_item_id=grn_item.id, qty=10, cost_rate=80, is_void=False))
        db.session.add(SupplierPayment(supplier_id=supplier.id, amount=500, date_posted=datetime(2026, 1, 6), is_void=False))
        db.session.add(AccountTransaction(amount=500, transaction_type="Payment", source_type="SupplierPayment", source_id=1, date_posted=datetime(2026, 1, 6), is_void=False))
        db.session.commit()

        result = report("2026-01-01", "2026-01-31", include_position=True)
        assert result.summary["total_revenue"] == pytest.approx(1200)
        assert result.summary["total_cogs"] == pytest.approx(800)
        assert result.summary["total_profit"] == pytest.approx(400)
        assert result.summary["total_expenses"] == 0
        assert result.rows[0]["cost_status"] == "actual"
        assert result.financial_position["purchase_value"] == pytest.approx(8000)
        before_supply = report("2026-01-01", "2026-01-04")
        assert before_supply.summary["total_profit"] == 0


def test_03_booking_partial_standalone_supply_and_global_fifo(pl_app):
    with pl_app.app_context():
        material, client, supplier, _account = seed_base(material_rate=999)
        _grn1, lot1 = add_grn(supplier, material, 50, 80, datetime(2026, 1, 1), "GRN-80")
        _grn2, lot2 = add_grn(supplier, material, 50, 90, datetime(2026, 1, 2), "GRN-90")
        booking, booked = add_booking(client, datetime(2026, 1, 3), [("CEMENT", 100, 100)], paid=5000)
        db.session.commit()

        booking_only = report("2026-01-03", "2026-01-31", include_position=True)
        assert booking_only.summary["total_revenue"] == 0
        assert booking_only.summary["total_profit"] == 0
        assert booking_only.financial_position["outstanding_booking_value"] == pytest.approx(10000)
        assert booking_only.financial_position["customer_advances_estimate"] == pytest.approx(5000)

        supply, supply_items = add_sale(client, datetime(2026, 2, 10), [("CEMENT", 40, 0, None)], category="Booking Delivery", bill="SUPPLY-40")
        db.session.add(BookingAllocation(sale_id=supply.id, sale_item_id=supply_items[0].id, booking_item_id=booked[0].id, qty=40, is_void=False))
        # Twenty chargeable units consume the rest of lot 1 and ten of lot 2
        # between the two booked supplies.  The second booked supply must see
        # that consumption instead of reusing a per-material latest rate.
        direct, direct_items = add_sale(client, datetime(2026, 2, 12), [("CEMENT", 20, 130, None)], bill="DIRECT-20")
        db.session.add_all([
            GRNAllocation(sale_id=direct.id, sale_item_id=direct_items[0].id, grn_item_id=lot1.id, qty=10, cost_rate=80, is_void=False),
            GRNAllocation(sale_id=direct.id, sale_item_id=direct_items[0].id, grn_item_id=lot2.id, qty=10, cost_rate=90, is_void=False),
        ])
        standalone = Entry(date="2026-02-15", time="10:00:00", type="OUT", material="CEMENT", booked_material="CEMENT", client=client.name, client_code=client.code, qty=10, bill_no="BK-1", transaction_category="Booked", client_category="Booking Delivery", nimbus_no="Booking Delivery", is_void=False)
        unmatched = Entry(date="2026-02-16", time="10:00:00", type="OUT", material="UNBOOKED", client=client.name, client_code=client.code, qty=5, transaction_category="Booked", is_void=False)
        db.session.add_all([standalone, unmatched])
        db.session.commit()

        first = report("2026-02-10", "2026-02-10")
        assert first.summary["total_revenue"] == pytest.approx(4000)
        assert first.summary["total_cogs"] == pytest.approx(3200)
        assert first.summary["total_profit"] == pytest.approx(800)
        assert first.rows[0]["cost_status"] == "estimated"
        assert "global FIFO" in first.rows[0]["cost_basis"]

        direct_day = report("2026-02-12", "2026-02-12")
        assert direct_day.summary["total_profit"] == pytest.approx(900)
        standalone_day = report("2026-02-15", "2026-02-15")
        assert standalone_day.summary["total_revenue"] == pytest.approx(1000)
        assert standalone_day.summary["total_cogs"] == pytest.approx(900)
        assert standalone_day.summary["total_profit"] == pytest.approx(100)
        assert standalone_day.rows[0]["source"] == "Entry"

        full = report("2026-02-01", "2026-02-29", include_position=True)
        assert full.summary["total_revenue"] == pytest.approx(7600)
        assert full.summary["total_cogs"] == pytest.approx(5800)
        assert full.summary["total_profit"] == pytest.approx(1800)
        assert full.financial_position["outstanding_booking_value"] == pytest.approx(5000)
        assert any(w["code"] == "unmatched_dispatch" for w in full.warnings)
        assert booking.id and lot1.id and lot2.id and direct.id


def test_04_booking_discount_is_prorated_to_actual_supply(pl_app):
    with pl_app.app_context():
        _material, client, _supplier, _account = seed_base(material_rate=50)
        booking, booked = add_booking(client, datetime(2026, 1, 1), [("CEMENT", 10, 100)], discount=100)
        sale, sale_items = add_sale(client, datetime(2026, 1, 5), [("CEMENT", 4, 0, 50)], category="Booking Delivery")
        db.session.add(BookingAllocation(sale_id=sale.id, sale_item_id=sale_items[0].id, booking_item_id=booked[0].id, qty=4, is_void=False))
        db.session.commit()
        result = report("2026-01-01", "2026-01-31", include_position=True)
        assert result.summary["total_revenue"] == pytest.approx(360)
        assert result.summary["total_discounts"] == pytest.approx(40)
        assert result.summary["total_cogs"] == pytest.approx(200)
        assert result.summary["total_profit"] == pytest.approx(160)
        assert result.financial_position["outstanding_booking_value"] == pytest.approx(540)
        booked[0].price_at_time = 120
        db.session.commit()
        edited = report("2026-01-01", "2026-01-31")
        assert edited.summary["total_revenue"] == pytest.approx(440)
        assert any("rate-history" in note for note in edited.unsupported)
        assert booking.id


def test_05_alternate_booked_material_and_excess_are_not_duplicated(pl_app):
    with pl_app.app_context():
        original = Material(name="12MM STEEL", code="ST-12", unit="Kg", unit_price=200, total=1000)
        delivered = Material(name="20MM STEEL", code="ST-20", unit="Kg", unit_price=210, total=2000)
        client = Client(name="CUSTOMER ONE", code="CL-1", category="General", is_active=True)
        supplier = Supplier(name="SUPPLIER ONE", is_active=True)
        db.session.add_all([original, delivered, client, supplier])
        db.session.flush()
        _grn, grn_item = add_grn(supplier, delivered, 2000, 210, datetime(2026, 1, 1))
        booking, booked = add_booking(client, datetime(2026, 1, 2), [("12MM STEEL", 1000, 240)])
        sale, items = add_sale(client, datetime(2026, 1, 5), [("20MM STEEL", 1000, 0, None), ("20MM STEEL", 100, 250, 210)], category="Mixed Transaction", bill="ALT-1")
        items[1].grn_item_id = grn_item.id
        db.session.add(BookingAllocation(sale_id=sale.id, sale_item_id=items[0].id, booking_item_id=booked[0].id, qty=1000, is_void=False))
        db.session.add(GRNAllocation(sale_id=sale.id, sale_item_id=items[1].id, grn_item_id=grn_item.id, qty=100, cost_rate=210, is_void=False))
        db.session.commit()
        result = report("2026-01-01", "2026-01-31")
        assert result.summary["total_revenue"] == pytest.approx(265000)
        assert result.summary["total_cogs"] == pytest.approx(231000)
        assert result.summary["total_profit"] == pytest.approx(34000)
        assert len(result.rows) == 2
        assert {row["material"] for row in result.rows} == {"20MM STEEL"}
        booked_row = next(row for row in result.rows if row["entry_type"] == "booking_supply")
        excess_row = next(row for row in result.rows if row["entry_type"] == "material_sale")
        assert booked_row["booked_material"] == "12MM STEEL"
        assert booked_row["qty"] == pytest.approx(1000)
        assert excess_row["qty"] == pytest.approx(100)
        assert sum(row["qty"] for row in result.rows) == pytest.approx(1100)
        assert booking.id


def test_06_recorded_grn_and_exact_fifo_lot_rate_beat_frozen_and_master(pl_app):
    with pl_app.app_context():
        material, client, supplier, _account = seed_base(material_rate=999)
        _grn, grn_item = add_grn(supplier, material, 10, 80, datetime(2026, 1, 1))
        sale, items = add_sale(client, datetime(2026, 1, 5), [("CEMENT", 5, 150, 85), ("CEMENT", 2, 150, 85)])
        items[0].grn_item_id = grn_item.id
        db.session.add(GRNAllocation(sale_id=sale.id, sale_item_id=items[0].id, grn_item_id=grn_item.id, qty=5, cost_rate=80, is_void=False))
        db.session.commit()
        material.unit_price = 1200
        db.session.commit()
        result = report("2026-01-01", "2026-01-31")
        assert result.summary["total_revenue"] == pytest.approx(1050)
        # The second line's frozen 85 is not mixed over a real 80-rate lot:
        # 5 x 80 recorded + 2 x 80 reconstructed = 560.
        assert result.summary["total_cogs"] == pytest.approx(560)
        assert result.summary["actual_material_cost"] == pytest.approx(400)
        assert result.summary["estimated_material_cost"] == pytest.approx(160)
        assert result.summary["total_profit"] == pytest.approx(490)
        assert [row["cost_status"] for row in result.rows] == ["actual", "estimated"]
        assert all(row["cost_sources"] for row in result.rows)
        assert all(
            source["grn_item_id"] == grn_item.id
            for row in result.rows for source in row["cost_sources"]
        )


def test_07_missing_cost_is_not_zero_and_estimate_is_explicit(pl_app):
    with pl_app.app_context():
        known, client, _supplier, _account = seed_base(material_rate=90)
        missing = Material(name="NO COST ITEM", code="NC-1", unit="Pcs", unit_price=0, total=10)
        db.session.add(missing)
        db.session.flush()
        add_sale(client, datetime(2026, 1, 5), [("NO COST ITEM", 2, 150, None)])
        add_sale(client, datetime(2026, 1, 6), [("CEMENT", 1, 150, None)])
        db.session.commit()
        result = report("2026-01-01", "2026-01-31")
        assert result.summary["total_revenue"] == pytest.approx(450)
        assert result.summary["total_cogs"] == pytest.approx(90)
        assert result.summary["total_profit"] == pytest.approx(60)
        assert result.summary["unresolved_revenue"] == pytest.approx(300)
        assert result.summary["unknown_cost_rows"] == 1
        assert result.summary["estimated_cost_rows"] == 1
        assert result.summary["formula_difference"] == pytest.approx(0)
        missing_row = next(row for row in result.rows if row["material"] == "NO COST ITEM")
        assert not missing_row["include_in_net"]
        assert missing_row["profit"] == 0
        assert known.id


def test_08_normal_return_reverses_original_revenue_and_cost_on_return_date(pl_app):
    with pl_app.app_context():
        _material, client, _supplier, _account = seed_base(material_rate=100)
        sale, _items = add_sale(client, datetime(2026, 1, 10), [("CEMENT", 10, 150, 100)], bill="SALE-RETURN")
        ret = MaterialReturn(client_name=client.name, return_type="normal", amount=600, manual_bill_no="RET-1", date_posted=datetime(2026, 1, 15), is_void=False)
        db.session.add(ret)
        db.session.flush()
        db.session.add(MaterialReturnItem(material_return_id=ret.id, material_name="CEMENT", qty=4, unit_rate=150, price_at_time=150))
        db.session.add(Payment(client_id=client.id, client_name=client.name, amount=600, payment_type="Material Return", source_type="MaterialReturn", source_id=ret.id, date_posted=datetime(2026, 1, 15), is_void=False))
        db.session.commit()
        before = report("2026-01-10", "2026-01-14")
        assert before.summary["total_profit"] == pytest.approx(500)
        return_day = report("2026-01-15", "2026-01-15")
        assert return_day.summary["total_revenue"] == pytest.approx(-600)
        assert return_day.summary["total_cogs"] == pytest.approx(-400)
        assert return_day.summary["total_profit"] == pytest.approx(-200)
        assert len(return_day.rows) == 1
        assert return_day.rows[0]["original_sources"][0]["id"] == sale.id
        full = report("2026-01-01", "2026-01-31")
        assert full.summary["total_revenue"] == pytest.approx(900)
        assert full.summary["total_cogs"] == pytest.approx(600)
        assert full.summary["total_profit"] == pytest.approx(300)
        assert len(report("2026-01-01", "2026-01-31").rows) == len(full.rows)

        # The four returned units form a dated cost layer.  A later sale must
        # consume that layer, not the subsequently changed master price.
        material = Material.query.filter_by(name="CEMENT").first()
        material.unit_price = 999
        add_sale(client, datetime(2026, 1, 20), [("CEMENT", 2, 160, None)], bill="AFTER-RETURN")
        db.session.commit()
        later = report("2026-01-20", "2026-01-20")
        assert later.summary["total_cogs"] == pytest.approx(200)
        assert later.summary["total_profit"] == pytest.approx(120)
        assert any(source["type"] == "ReturnLot" for source in later.rows[0]["cost_sources"])


def test_09_booked_return_reopens_unfulfilled_quantity_without_touching_advance_profit(pl_app):
    with pl_app.app_context():
        _material, client, _supplier, _account = seed_base(material_rate=80)
        booking, booked = add_booking(client, datetime(2026, 1, 1), [("CEMENT", 10, 120)], paid=1200)
        sale, sale_items = add_sale(client, datetime(2026, 2, 1), [("CEMENT", 10, 0, 80)], category="Booking Delivery")
        db.session.add(BookingAllocation(sale_id=sale.id, sale_item_id=sale_items[0].id, booking_item_id=booked[0].id, qty=10, is_void=False))
        ret = MaterialReturn(client_name=client.name, return_type="booked", amount=360, manual_bill_no="BRET-1", date_posted=datetime(2026, 2, 5), is_void=False)
        db.session.add(ret)
        db.session.flush()
        db.session.add(MaterialReturnItem(material_return_id=ret.id, material_name="CEMENT", qty=3, rent_rate=120, price_at_time=120))
        db.session.commit()
        before = report("2026-02-01", "2026-02-04")
        assert before.summary["total_profit"] == pytest.approx(400)
        return_day = report("2026-02-05", "2026-02-05")
        assert return_day.summary["total_profit"] == pytest.approx(-120)
        full = report("2026-02-01", "2026-02-28", include_position=True)
        assert full.summary["total_revenue"] == pytest.approx(840)
        assert full.summary["total_cogs"] == pytest.approx(560)
        assert full.summary["total_profit"] == pytest.approx(280)
        assert full.financial_position["outstanding_booking_value"] == pytest.approx(360)
        assert full.financial_position["customer_advances_estimate"] == pytest.approx(360)
        assert booking.id


def test_10_delivery_profit_is_once_and_driver_payment_is_settlement(pl_app):
    with pl_app.app_context():
        _material, client, _supplier, _account = seed_base()
        db.session.add(Service(name="RENT-CEMENT", code="SVC-RENT", unit="Bags"))
        driver = DeliveryPerson(name="DRIVER ONE", is_active=True)
        db.session.add(driver)
        db.session.flush()
        sale, _items = add_sale(client, datetime(2026, 2, 1), [("RENT-CEMENT", 10, 50, None)], delivery_cost=300, bill="RENT-1")
        db.session.add(DeliveryPersonPayment(delivery_person_id=driver.id, sale_id=sale.id, amount_paid=300, waive_off_amount=0, date_posted=datetime(2026, 2, 2), is_void=False))
        db.session.add(AccountTransaction(amount=300, transaction_type="Driver Payment", source_type="DeliveryPersonPayment", source_id=1, date_posted=datetime(2026, 2, 2), is_void=False))
        db.session.commit()
        first = report("2026-02-01", "2026-02-01")
        assert first.summary["total_delivery_revenue"] == pytest.approx(500)
        assert first.summary["total_delivery_cost"] == pytest.approx(300)
        assert first.summary["total_profit"] == pytest.approx(200)
        assert len(first.rows) == 1
        full_before_waive = report("2026-02-01", "2026-02-28")
        assert full_before_waive.summary["total_profit"] == pytest.approx(200)
        db.session.add(DeliveryPersonPayment(delivery_person_id=driver.id, sale_id=sale.id, amount_paid=0, waive_off_amount=50, date_posted=datetime(2026, 2, 3), is_void=False))
        db.session.commit()
        full = report("2026-02-01", "2026-02-28")
        assert full.summary["total_other_gains"] == pytest.approx(50)
        assert full.summary["total_delivery_cost"] == pytest.approx(300)
        assert full.summary["total_profit"] == pytest.approx(250)
        assert report("2026-02-03", "2026-02-03").summary["total_profit"] == pytest.approx(50)

        # A free service can still have an actual delivery cost; zero revenue
        # must not make that cost disappear from the day or from net profit.
        add_sale(client, datetime(2026, 2, 4), [("RENT-CEMENT", 2, 0, None)], delivery_cost=100, bill="FREE-RENT-1")
        db.session.commit()
        free_delivery = report("2026-02-04", "2026-02-04")
        assert free_delivery.summary["total_delivery_revenue"] == 0
        assert free_delivery.summary["total_delivery_cost"] == pytest.approx(100)
        assert free_delivery.summary["total_profit"] == pytest.approx(-100)

        # Legacy records may carry the earned rent only on the sale header.
        legacy = DirectSale(client_name=client.name, client_code=client.code, category="Credit Customer", amount=200, rent_item_revenue=200, delivery_rent_cost=50, manual_bill_no="LEGACY-RENT-1", date_posted=datetime(2026, 2, 5), is_void=False)
        db.session.add(legacy)
        db.session.commit()
        legacy_day = report("2026-02-05", "2026-02-05")
        assert legacy_day.summary["total_delivery_revenue"] == pytest.approx(200)
        assert legacy_day.summary["total_delivery_cost"] == pytest.approx(50)
        assert legacy_day.summary["total_profit"] == pytest.approx(150)


def test_11_expenses_reconciliation_and_other_income_are_counted_once(pl_app):
    with pl_app.app_context():
        _material, _client, _supplier, account = seed_base()
        expense = AccountTransaction(from_account_id=account.id, amount=100, description="Fuel", transaction_type="Expense", source_type="CashFlowEntry", source_id=1, date_posted=datetime(2026, 3, 1), is_void=False)
        db.session.add(expense)
        db.session.flush()
        db.session.add(CashFlowEntry(direction="out", amount=100, account_id=account.id, account_tx_id=expense.id, description="Fuel", date_posted=datetime(2026, 3, 1), is_void=False))
        db.session.add_all([
            FbmCashDrawerEntry(entry_type="out", amount=40, category="Shop expense", date_posted=datetime(2026, 3, 2), is_void=False),
            AccountTransaction(from_account_id=account.id, amount=20, transaction_type="Reconciliation Loss", date_posted=datetime(2026, 3, 3), is_void=False),
            AccountTransaction(to_account_id=account.id, amount=5, transaction_type="Reconciliation Excess", date_posted=datetime(2026, 3, 4), is_void=False),
            AccountTransaction(to_account_id=account.id, amount=70, description="Money received from Scrap", transaction_type="Receipt", date_posted=datetime(2026, 3, 5), is_void=False),
        ])
        db.session.commit()
        result = report("2026-03-01", "2026-03-31")
        assert result.summary["total_expenses"] == pytest.approx(140)
        assert result.summary["total_other_losses"] == pytest.approx(20)
        assert result.summary["total_other_gains"] == pytest.approx(5)
        assert result.summary["total_other_income"] == pytest.approx(70)
        assert result.summary["total_revenue"] == pytest.approx(70)
        assert result.summary["total_profit"] == pytest.approx(-85)
        assert len(result.rows) == 5


def test_12_discounts_waives_and_voids_are_not_double_counted(pl_app):
    with pl_app.app_context():
        _material, client, _supplier, _account = seed_base()
        sale, _items = add_sale(client, datetime(2026, 1, 1), [("CEMENT", 10, 100, 50)], discount=100, bill="DISCOUNT-1")
        payment = Payment(client_id=client.id, client_name=client.name, amount=900, discount=50, payment_type="Receipt", manual_bill_no="PAY-1", date_posted=datetime(2026, 1, 3), is_void=False)
        legacy = Payment(client_id=client.id, client_name=client.name, amount=0, discount=30, payment_type="Waive-Off", manual_bill_no="PAY-2", date_posted=datetime(2026, 1, 4), is_void=False)
        db.session.add_all([payment, legacy])
        db.session.flush()
        db.session.add_all([
            WaiveOff(note=f"[direct_sale_discount:{sale.id}]", client_name=client.name, bill_no="DISCOUNT-1", amount=100, date_posted=datetime(2026, 1, 1), is_void=False),
            WaiveOff(payment_id=payment.id, client_name=client.name, bill_no="PAY-1", amount=50, date_posted=datetime(2026, 1, 3), is_void=False),
            WaiveOff(payment_id=payment.id, client_name=client.name, bill_no="PAY-1", amount=25, date_posted=datetime(2026, 1, 3), is_void=True),
        ])
        add_sale(client, datetime(2026, 1, 5), [("CEMENT", 5, 200, 100)], void=True)
        db.session.commit()
        result = report("2026-01-01", "2026-01-31")
        assert result.summary["total_revenue"] == pytest.approx(900)
        assert result.summary["total_discounts"] == pytest.approx(100)
        assert result.summary["total_other_losses"] == pytest.approx(80)
        assert result.summary["total_cogs"] == pytest.approx(500)
        assert result.summary["total_profit"] == pytest.approx(320)
        assert not any(row["source"] == "WaiveOff" and row["other_loss"] == 100 for row in result.rows)
        assert not any(row["document_type"] == "DirectSale" and row["document_id"] != sale.id for row in result.rows)
        assert result.voided_source_count >= 1

        add_sale(client, datetime(2026, 1, 6), [("CEMENT", 1, 0, 25)], bill="FREE-MATERIAL-1")
        db.session.commit()
        free_material = report("2026-01-06", "2026-01-06")
        assert free_material.summary["total_revenue"] == 0
        assert free_material.summary["total_cogs"] == pytest.approx(25)
        assert free_material.summary["total_profit"] == pytest.approx(-25)
        assert free_material.rows[0]["cost_status"] == "actual"


def test_13_prior_period_filters_pagination_and_presets(pl_app):
    with pl_app.app_context():
        _material, client, _supplier, account = seed_base()
        other = Client(name="OTHER CUSTOMER", code="CL-2", category="General", is_active=True)
        db.session.add(other)
        db.session.flush()
        add_sale(client, datetime(2026, 1, 10), [("CEMENT", 1, 150, 50)])
        add_sale(client, datetime(2026, 2, 10), [("CEMENT", 1, 200, 50)])
        add_sale(other, datetime(2026, 2, 11), [("CEMENT", 1, 500, 100)])
        add_sale(client, datetime(2026, 2, 12), [("CEMENT", 10, 1000, 1)], void=True)
        for index in range(30):
            db.session.add(AccountTransaction(from_account_id=account.id, amount=1, description=f"Expense {index}", transaction_type="Expense", date_posted=datetime(2026, 2, 15), is_void=False))
        db.session.commit()
        result = report("2026-02-01", "2026-02-28", client="CL-1")
        assert result.summary["total_profit"] == pytest.approx(150)
        assert result.previous_summary["total_profit"] == pytest.approx(100)
        assert result.comparison["metrics"]["total_profit"]["change"] == pytest.approx(50)
        assert result.comparison["label"] == "Previous full month"
        paged = report("2026-02-01", "2026-02-28", page=2, per_page=10)
        assert paged.entries_count == 32
        assert paged.pagination["pages"] == 4
        assert len(paged.transactions) == 10
        material_only = report("2026-02-01", "2026-02-28", material="CEMENT")
        assert material_only.entries_count == 2
        start, end, _label = resolve_report_period("last_month", today=date(2026, 3, 15))
        assert (start, end) == (date(2026, 2, 1), date(2026, 2, 28))


def test_14_dashboard_detail_pdf_excel_reconcile_and_permission(pl_app):
    with pl_app.app_context():
        _material, client, _supplier, account = seed_base()
        missing = Material(name="NO COST ITEM", code="NC-1", unit="Pcs", unit_price=0, total=10)
        db.session.add(missing)
        db.session.flush()
        sale, _items = add_sale(client, datetime(2026, 4, 5), [("CEMENT", 2, 150, 100)], bill="EXPORT-1")
        add_sale(client, datetime(2026, 4, 6), [("NO COST ITEM", 1, 500, None)], bill="EXPORT-MISSING")
        db.session.add(AccountTransaction(from_account_id=account.id, amount=25, description="Office", transaction_type="Expense", date_posted=datetime(2026, 4, 7), is_void=False))
        db.session.add(User(username="noreports", password_hash=generate_password_hash("NoReports123!"), role="user", status="active", can_view_reports=False))
        db.session.commit()
        expected = report("2026-04-01", "2026-04-30")
        assert expected.summary["total_profit"] == pytest.approx(75)
        assert expected.summary["total_revenue"] == pytest.approx(800)
        detail = build_transaction_detail("DirectSale", sale.id, start_date="2026-04-01", end_date="2026-04-30")
        assert detail["summary"]["total_profit"] == pytest.approx(100)

    browser = pl_app.test_client()
    login(browser)
    dashboard = browser.get("/profit_reports?start_date=2026-04-01&end_date=2026-04-30")
    assert dashboard.status_code == 200
    assert b"EXPORT-1" in dashboard.data
    detail_response = browser.get(f"/profit_reports/transaction/DirectSale/{sale.id}?start_date=2026-04-01&end_date=2026-04-30")
    assert detail_response.status_code == 200
    assert b"Lifetime P&L Components" in detail_response.data
    assert b"EXPORT-1" in detail_response.data

    excel_response = browser.get("/profit_reports/export/excel?start_date=2026-04-01&end_date=2026-04-30")
    assert excel_response.status_code == 200
    assert excel_response.data[:2] == b"PK"
    workbook = load_workbook(BytesIO(excel_response.data), data_only=True)
    summary_ws = workbook["Summary"]
    labels = {summary_ws.cell(row, 1).value: row for row in range(1, summary_ws.max_row + 1)}
    assert summary_ws.cell(labels["Net Profit / Loss"], 2).value == pytest.approx(75)
    assert summary_ws.cell(labels["Sales (Money In)"], 2).value == pytest.approx(800)
    tx_ws = workbook["Transactions"]
    headers = {cell.value: cell.column for cell in tx_ws[1]}
    assert sum(float(tx_ws.cell(row, headers["Net Profit"]).value or 0) for row in range(2, tx_ws.max_row + 1)) == pytest.approx(75)
    assert sum(float(tx_ws.cell(row, headers["Material Revenue"]).value or 0) + float(tx_ws.cell(row, headers["Delivery Revenue"]).value or 0) + float(tx_ws.cell(row, headers["Other Income"]).value or 0) for row in range(2, tx_ws.max_row + 1)) == pytest.approx(800)

    pdf_response = browser.get("/profit_reports/export/pdf?start_date=2026-04-01&end_date=2026-04-30")
    assert pdf_response.status_code == 200
    assert pdf_response.data.startswith(b"%PDF-")
    pdf_text = "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(pdf_response.data)).pages)
    assert "Net Profit / Loss" in pdf_text
    assert "75.00" in pdf_text
    assert "EXPORT-1" in "".join(pdf_text.split())

    denied = pl_app.test_client()
    login(denied, "noreports", "NoReports123!")
    for path in (
        "/profit_reports",
        f"/profit_reports/transaction/DirectSale/{sale.id}",
        "/profit_reports/export/pdf",
        "/profit_reports/export/excel",
    ):
        response = denied.get(path, follow_redirects=False)
        assert response.status_code in (302, 303)
        assert response.headers["Location"].endswith("/")
