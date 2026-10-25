"""Multi-GRN cost of goods sold in the Profit & Loss report.

The required scenario: 10 bags bought at Rs. 1,000 (old GRN) and 30 bags at
Rs. 1,200 (new GRN), all 40 delivered.  Cost of goods sold must be
10 x 1,000 + 30 x 1,200 = Rs. 46,000 (weighted average Rs. 1,150/bag) - never
40 x 1,200 = 48,000, 40 x 1,000 = 40,000, or any other single GRN rate.

These tests pin the whole chain: GRN receipt -> FIFO lot allocation -> delivery
-> cost reconstruction -> report rows -> dashboard/detail/Excel exports, plus
returns, partial deliveries and cancellations.  Every assertion is computed
from the isolated test database created per test; nothing touches live data.
"""
from __future__ import annotations

from datetime import datetime
from io import BytesIO
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from app.services.grn_consumption import build_grn_consumption
from app.services.profit_loss import build_profit_loss_report
from models import (
    Client,
    DirectSale,
    DirectSaleItem,
    GRN,
    GRNAllocation,
    GRNItem,
    Material,
    Supplier,
    db,
)
from tests.conftest import make_csrf_client

EXPECTED_COGS = 10 * 1000 + 30 * 1200      # Rs. 46,000
EXPECTED_AVG = EXPECTED_COGS / 40          # Rs. 1,150 per bag


@pytest.fixture()
def multi_app(app_factory, tmp_path):
    app = app_factory(tmp_path / "multi_grn.db")
    app.config["TESTING"] = True
    with app.app_context():
        db.create_all()
        from app.services.schema import _ensure_default_admin
        _ensure_default_admin()
        db.session.commit()
    return app


def login(client):
    response = client.post("/login", data={"username": "Admin", "password": "Admin@fbm12345"})
    assert response.status_code in (302, 303)


def seed(app):
    with app.app_context():
        material = Material(name="CEMENT", code="MAT-1", unit="Bags", unit_price=999, total=0)
        client_row = Client(name="CUSTOMER ONE", code="CL-1", category="General", is_active=True)
        supplier = Supplier(name="SUPPLIER ONE", is_active=True)
        db.session.add_all([material, client_row, supplier])
        db.session.commit()
        return material.id, client_row.id, supplier.id


def add_grn(app, supplier_id, qty, rate, posted, bill):
    with app.app_context():
        supplier = db.session.get(Supplier, supplier_id)
        grn = GRN(
            supplier_id=supplier.id, supplier=supplier.name, manual_bill_no=bill,
            date_posted=posted, is_void=False,
        )
        db.session.add(grn)
        db.session.flush()
        item = GRNItem(grn_id=grn.id, mat_name="CEMENT", qty=qty, price_at_time=rate)
        db.session.add(item)
        db.session.flush()
        grn_id, item_id = grn.id, item.id
        db.session.commit()
        return grn_id, item_id


def post_delivery(client, qty, rate=1300, bill="SALE-1", sale_date="2026-01-10"):
    response = client.post("/add_direct_sale", data={
        "client_name": "CUSTOMER ONE",
        "driver_name": "Driver",
        "category": "Credit Customer",
        "product_name[]": ["CEMENT"],
        "qty[]": [str(qty)],
        "unit_rate[]": [str(rate)],
        "payment_method": "Credit",
        "paid_amount": "0",
        "manual_bill_no": bill,
        "sale_date": sale_date,
        "allow_negative_stock": "1",
    }, follow_redirects=False)
    assert response.status_code in (302, 303), response.get_data(as_text=True)[:400]


def report(start="2026-01-01", end="2026-12-31", **kwargs):
    return SimpleNamespace(**build_profit_loss_report(start_date=start, end_date=end, **kwargs))


def goods_rows(rep, entry_type=("material_sale", "booking_supply")):
    return [row for row in rep.rows if row["entry_type"] in entry_type and row["material"] == "CEMENT"]


def posted_sale(app, bill):
    with app.app_context():
        return DirectSale.query.filter(DirectSale.manual_bill_no.like(f"%{bill}%")).one().id


# --------------------------------------------------------------------------- #
# 1. The exact required case
# --------------------------------------------------------------------------- #
def test_required_10_plus_30_bag_delivery_costs_rs_46000(multi_app):
    client = make_csrf_client(multi_app)
    login(client)
    with multi_app.app_context():
        _material, _client_row, supplier_id = seed(multi_app)
        add_grn(multi_app, supplier_id, 10, 1000, datetime(2026, 1, 1, 9, 0), "GRN-OLD")
        add_grn(multi_app, supplier_id, 30, 1200, datetime(2026, 1, 5, 9, 0), "GRN-NEW")
        post_delivery(client, 40, rate=1300, bill="TEN-THIRTY")
        sale_id = posted_sale(multi_app, "TEN-THIRTY")

        # The posting path records one sale line per FIFO lot, each with its own
        # frozen purchase rate: this is the recorded source allocation.
        with multi_app.app_context():
            items = DirectSaleItem.query.filter_by(sale_id=sale_id).order_by(DirectSaleItem.id).all()
            assert [(item.qty, item.cost_rate_at_sale) for item in items] == [(10.0, 1000.0), (30.0, 1200.0)]
            allocations = GRNAllocation.query.filter_by(sale_id=sale_id).order_by(GRNAllocation.id).all()
            assert [round(a.qty, 2) for a in allocations] == [10.0, 30.0]
            assert [a.cost_rate for a in allocations] == [1000.0, 1200.0]
            assert allocations[0].grn_item_id != allocations[1].grn_item_id
            # Selling price is untouched by the purchase rates.
            assert [item.price_at_time for item in items] == [1300.0, 1300.0]

        rep = report()
        rows = goods_rows(rep)
        assert [round(row["qty"], 2) for row in rows] == [10.0, 30.0]
        assert [row["material_cost"] for row in rows] == [10000.0, 36000.0]
        assert [row["cost_rate"] for row in rows] == [1000.0, 1200.0]
        assert all(row["cost_status"] == "actual" for row in rows)
        assert rep.summary["total_cogs"] == pytest.approx(EXPECTED_COGS)
        assert rep.summary["total_cogs"] != pytest.approx(48_000)      # not 40 x 1,200
        assert rep.summary["total_cogs"] != pytest.approx(40_000)      # not 40 x 1,000
        assert rep.summary["total_material_revenue"] == pytest.approx(52_000)
        assert rep.summary["total_profit"] == pytest.approx(52_000 - EXPECTED_COGS)
        assert sum(row["qty"] for row in rows) == pytest.approx(40)
        assert sum(row["material_cost"] for row in rows) / 40 == pytest.approx(EXPECTED_AVG)

        # Both GRN lots appear on the delivery with their own quantity and rate.
        for row in rows:
            assert len(row["cost_sources"]) == 1
            source = row["cost_sources"][0]
            assert source["reference"] in {"MB NO.GRN-OLD", "MB NO.GRN-NEW", "GRN-OLD", "GRN-NEW"}
            assert source["qty"] == pytest.approx(row["qty"])
            assert source["amount"] == pytest.approx(row["material_cost"])
        references = {row["cost_sources"][0]["reference"] for row in rows}
        assert len(references) == 2, f"both source GRNs must be shown, got {references}"

        # Every lot is fully consumed and nothing is left behind.
        consumption = build_grn_consumption()
        lots = sorted(consumption["items"].values(), key=lambda lot: lot["received_qty"])
        assert [(lot["received_qty"], lot["consumed_qty"], lot["remaining_qty"]) for lot in lots] == [
            (10.0, 10.0, 0.0), (30.0, 30.0, 0.0),
        ]
        assert consumption["inventory"]["total_qty"] == pytest.approx(0.0)
        assert rep.context.grn_lot_balances
        assert all(lot["remaining"] == pytest.approx(0.0) for lot in rep.context.grn_lot_balances.values())


# --------------------------------------------------------------------------- #
# 2. Transaction detail, dashboard, exports and repeatability
# --------------------------------------------------------------------------- #
def test_detail_dashboard_and_excel_exports_reconcile(multi_app):
    client = make_csrf_client(multi_app)
    login(client)
    with multi_app.app_context():
        _material, _client_row, supplier_id = seed(multi_app)
        add_grn(multi_app, supplier_id, 10, 1000, datetime(2026, 1, 1, 9, 0), "GRN-OLD")
        add_grn(multi_app, supplier_id, 30, 1200, datetime(2026, 1, 5, 9, 0), "GRN-NEW")
        post_delivery(client, 40, rate=1300, bill="EXPORT-46K")
        sale_id = posted_sale(multi_app, "EXPORT-46K")

        # Repeated report generation must not double count or drift.
        first = report()
        second = report()
        assert first.summary["total_cogs"] == second.summary["total_cogs"] == pytest.approx(EXPECTED_COGS)
        assert [row["material_cost"] for row in goods_rows(first)] == [row["material_cost"] for row in goods_rows(second)]

        # Dashboard page shows the same authoritative total.
        dashboard = client.get("/profit_reports?start_date=2026-01-01&end_date=2026-12-31")
        assert dashboard.status_code == 200
        assert b"EXPORT-46K" in dashboard.data

        # Transaction detail: delivered qty, both source GRNs, per-lot qty x rate.
        detail = client.get(f"/profit_reports/transaction/DirectSale/{sale_id}?start_date=2026-01-01&end_date=2026-12-31")
        assert detail.status_code == 200
        assert b"GRN cost layers (qty \xc3\x97 unit cost)" in detail.data
        assert b"GRN-OLD" in detail.data and b"GRN-NEW" in detail.data
        assert b"10.00 \xc3\x97 Rs. 1,000.00" in detail.data
        assert b"30.00 \xc3\x97 Rs. 1,200.00" in detail.data

        # Excel Transactions sheet totals to the same cost of goods sold.
        excel = client.get("/profit_reports/export/excel?start_date=2026-01-01&end_date=2026-12-31")
        assert excel.status_code == 200
        assert excel.data[:2] == b"PK"
        workbook = load_workbook(BytesIO(excel.data), data_only=True)
        tx_sheet = workbook["Transactions"]
        headers = {cell.value: cell.column for cell in tx_sheet[1]}
        material_cost_total = sum(
            float(tx_sheet.cell(row, headers["Material Cost"]).value or 0)
            for row in range(2, tx_sheet.max_row + 1)
        )
        assert material_cost_total == pytest.approx(EXPECTED_COGS)
        net_profit_total = sum(
            float(tx_sheet.cell(row, headers["Net Profit"]).value or 0)
            for row in range(2, tx_sheet.max_row + 1)
        )
        assert net_profit_total == pytest.approx(52_000 - EXPECTED_COGS)
        summary_sheet = workbook["Summary"]
        labels = {summary_sheet.cell(row, 1).value: row for row in range(1, summary_sheet.max_row + 1)}
        assert summary_sheet.cell(labels["Cost of Goods (Money Out)"], 2).value == pytest.approx(EXPECTED_COGS)

        # Every exported cost line carries its source GRN, quantity and rate.
        assert "Source GRNs" in headers and "GRN Cost Layers" in headers
        layers_by_grn = " ".join(
            str(tx_sheet.cell(row, headers["GRN Cost Layers"]).value or "")
            for row in range(2, tx_sheet.max_row + 1)
        )
        assert "GRN-OLD" in layers_by_grn and "GRN-NEW" in layers_by_grn
        assert "10.00 x Rs. 1,000.00" in layers_by_grn
        assert "30.00 x Rs. 1,200.00" in layers_by_grn
        source_grn_column = " ".join(
            str(tx_sheet.cell(row, headers["Source GRNs"]).value or "")
            for row in range(2, tx_sheet.max_row + 1)
        )
        assert "GRN-OLD" in source_grn_column and "GRN-NEW" in source_grn_column


# --------------------------------------------------------------------------- #
# 3. Delivery posted before the receipts (negative stock)
# --------------------------------------------------------------------------- #
def test_delivery_posted_before_receipts_is_valued_from_both_lots(multi_app):
    client = make_csrf_client(multi_app)
    login(client)
    with multi_app.app_context():
        _material, _client_row, supplier_id = seed(multi_app)
        # Nothing has been purchased when the 40 bags leave the yard.
        post_delivery(client, 40, rate=1300, bill="BEFORE-GRN", sale_date="2026-01-01")
        # No GRN exists yet: the only evidence is the material master price.
        pending = report()
        assert pending.summary["total_cogs"] == pytest.approx(40 * 999)
        assert goods_rows(pending)[0]["cost_status"] == "estimated"
        assert "no purchase history" in goods_rows(pending)[0]["cost_basis"]

        # The two purchases are recorded afterwards.
        add_grn(multi_app, supplier_id, 10, 1000, datetime(2026, 1, 5, 9, 0), "GRN-OLD")
        add_grn(multi_app, supplier_id, 30, 1200, datetime(2026, 1, 6, 9, 0), "GRN-NEW")

        rep = report()
        rows = goods_rows(rep)
        assert len(rows) == 1
        row = rows[0]
        assert row["qty"] == pytest.approx(40)
        assert row["material_cost"] == pytest.approx(EXPECTED_COGS)
        assert row["cost_rate"] == pytest.approx(EXPECTED_AVG)
        assert row["cost_status"] == "estimated"       # reconstruction, never "actual"
        assert "no purchase existed on/before supply" in row["cost_basis"]
        assert [(source["qty"], source["rate"]) for source in row["cost_sources"]] == [(10.0, 1000.0), (30.0, 1200.0)]
        assert {source["reference"] for source in row["cost_sources"]} == {"MB NO.GRN-OLD", "MB NO.GRN-NEW"}
        assert sum(source["amount"] for source in row["cost_sources"]) == pytest.approx(EXPECTED_COGS)
        assert rep.summary["total_cogs"] == pytest.approx(EXPECTED_COGS)
        # No allocation rows were invented for the historical delivery.
        with multi_app.app_context():
            assert GRNAllocation.query.count() == 0


# --------------------------------------------------------------------------- #
# 4. Partial deliveries
# --------------------------------------------------------------------------- #
def test_partial_deliveries_split_the_lots_correctly(multi_app):
    client = make_csrf_client(multi_app)
    login(client)
    with multi_app.app_context():
        _material, _client_row, supplier_id = seed(multi_app)
        add_grn(multi_app, supplier_id, 10, 1000, datetime(2026, 1, 1, 9, 0), "GRN-OLD")
        add_grn(multi_app, supplier_id, 30, 1200, datetime(2026, 1, 5, 9, 0), "GRN-NEW")
        post_delivery(client, 15, rate=1300, bill="PART-1")
        post_delivery(client, 25, rate=1300, bill="PART-2")

        rep = report()
        rows = goods_rows(rep)
        # 10 @ 1,000 + 5 @ 1,200 for the first 15 bags, then 25 @ 1,200.
        assert [(row["qty"], row["material_cost"]) for row in rows] == [
            (10.0, 10000.0), (5.0, 6000.0), (25.0, 30000.0),
        ]
        assert rep.summary["total_cogs"] == pytest.approx(EXPECTED_COGS)
        assert rep.summary["total_profit"] == pytest.approx(40 * 1300 - EXPECTED_COGS)
        consumption = build_grn_consumption()
        assert sum(lot["remaining_qty"] for lot in consumption["items"].values()) == pytest.approx(0.0)


# --------------------------------------------------------------------------- #
# 5. Returns reverse the original lots at their original unit cost
# --------------------------------------------------------------------------- #
def test_return_reverses_the_original_cost_layers(multi_app):
    client = make_csrf_client(multi_app)
    login(client)
    with multi_app.app_context():
        material_id, _client_row, supplier_id = seed(multi_app)
        add_grn(multi_app, supplier_id, 10, 1000, datetime(2026, 1, 1, 9, 0), "GRN-OLD")
        add_grn(multi_app, supplier_id, 30, 1200, datetime(2026, 1, 5, 9, 0), "GRN-NEW")
        post_delivery(client, 40, rate=1300, bill="RET-CASE")

        response = client.post("/add_material_return", data={
            "client_name": "CUSTOMER ONE",
            "date": "2026-01-12",
            "return_type": "normal",
            "material_name[]": ["CEMENT"],
            "material_id[]": [str(material_id)],
            "qty[]": ["10"],
            "unit_rate[]": ["1300"],
        }, follow_redirects=False)
        assert response.status_code in (302, 303), response.get_data(as_text=True)[:400]

        rep = report()
        returns = goods_rows(rep, ("return",))
        assert len(returns) == 1
        returned = returns[0]
        # The oldest lot (10 bags at Rs. 1,000) is reversed, not a blended rate.
        assert returned["qty"] == pytest.approx(-10)
        assert returned["material_cost"] == pytest.approx(-10000.0)
        assert returned["cost_rate"] == pytest.approx(1000.0)
        assert returned["cost_sources"]
        assert returned["cost_sources"][0]["reference"] == "MB NO.GRN-OLD"
        assert returned["cost_sources"][0]["qty"] == pytest.approx(-10)
        assert returned["cost_sources"][0]["amount"] == pytest.approx(-10000.0)
        assert rep.summary["total_cogs"] == pytest.approx(EXPECTED_COGS - 10000)
        assert rep.summary["total_profit"] == pytest.approx((40 - 10) * 1300 - (EXPECTED_COGS - 10000))


# --------------------------------------------------------------------------- #
# 6. Cancellation releases the lots and leaves no double count
# --------------------------------------------------------------------------- #
def test_cancelled_delivery_releases_its_lots(multi_app):
    client = make_csrf_client(multi_app)
    login(client)
    with multi_app.app_context():
        _material, _client_row, supplier_id = seed(multi_app)
        add_grn(multi_app, supplier_id, 10, 1000, datetime(2026, 1, 1, 9, 0), "GRN-OLD")
        add_grn(multi_app, supplier_id, 30, 1200, datetime(2026, 1, 5, 9, 0), "GRN-NEW")
        post_delivery(client, 40, rate=1300, bill="CANCEL-ME")
        assert report().summary["total_cogs"] == pytest.approx(EXPECTED_COGS)

        response = client.post("/delete_transaction/DirectSale/1", follow_redirects=False)
        assert response.status_code in (302, 303, 308)

        after = report()
        assert goods_rows(after) == []
        assert after.summary["total_cogs"] == pytest.approx(0.0)
        consumption = build_grn_consumption()
        assert [(lot["consumed_qty"], lot["remaining_qty"]) for lot in sorted(consumption["items"].values(), key=lambda lot: lot["received_qty"])] == [
            (0.0, 10.0), (0.0, 30.0),
        ]
        assert after.context.grn_lot_balances
        assert all(lot["remaining"] == pytest.approx(lot["received"]) for lot in after.context.grn_lot_balances.values())

        # The released stock can be delivered again, at the same lot rates.
        post_delivery(client, 40, rate=1300, bill="SECOND-TRY")
        again = report()
        assert again.summary["total_cogs"] == pytest.approx(EXPECTED_COGS)
        assert again.summary["total_profit"] == pytest.approx(52_000 - EXPECTED_COGS)


# --------------------------------------------------------------------------- #
# 7. A later purchase rate never changes the selling price or the old delivery
# --------------------------------------------------------------------------- #
def test_new_purchase_rate_does_not_rewrite_price_or_completed_delivery(multi_app):
    client = make_csrf_client(multi_app)
    login(client)
    with multi_app.app_context():
        _material, _client_row, supplier_id = seed(multi_app)
        add_grn(multi_app, supplier_id, 10, 1000, datetime(2026, 1, 1, 9, 0), "GRN-OLD")
        add_grn(multi_app, supplier_id, 30, 1200, datetime(2026, 1, 5, 9, 0), "GRN-NEW")
        post_delivery(client, 40, rate=1300, bill="STABLE")
        before = report()

        # A newer, more expensive purchase is recorded afterwards.
        add_grn(multi_app, supplier_id, 100, 1500, datetime(2026, 2, 1, 9, 0), "GRN-LATER")
        after = report()

        assert after.summary["total_cogs"] == pytest.approx(before.summary["total_cogs"]) == pytest.approx(EXPECTED_COGS)
        assert after.summary["total_material_revenue"] == pytest.approx(before.summary["total_material_revenue"])
        rows = goods_rows(after)
        assert [row["material_cost"] for row in rows] == [10000.0, 36000.0]
        with multi_app.app_context():
            items = DirectSaleItem.query.order_by(DirectSaleItem.id).all()
            assert [item.price_at_time for item in items] == [1300.0, 1300.0]
            sale = DirectSale.query.one()
            assert float(sale.amount or 0) == pytest.approx(52_000)


# --------------------------------------------------------------------------- #
# 8. An addendum GRN cannot re-cost a delivery that already had history
# --------------------------------------------------------------------------- #
def test_addendum_grn_does_not_re_cost_a_delivery_that_had_history(multi_app):
    client = make_csrf_client(multi_app)
    login(client)
    with multi_app.app_context():
        _material, _client_row, supplier_id = seed(multi_app)
        add_grn(multi_app, supplier_id, 10, 1000, datetime(2026, 1, 1, 9, 0), "GRN-OLD")
        # Only 10 bags are on hand, so posting freezes the known rate for the other
        # 30 bags and records no allocation for them.
        post_delivery(client, 40, rate=1300, bill="PARTIAL-STOCK")
        before = report()
        assert before.summary["total_cogs"] == pytest.approx(40 * 1000)

        # The supplier invoice for the remaining 30 bags is entered later, dated
        # before the delivery date.
        add_grn(multi_app, supplier_id, 30, 1200, datetime(2026, 1, 5, 9, 0), "GRN-ADDENDUM")
        after = report()

        # The recorded units keep their frozen rate; the addendum cannot silently
        # rewrite a completed delivery.
        assert after.summary["total_cogs"] == pytest.approx(40 * 1000)
        rows = goods_rows(after)
        assert [row["material_cost"] for row in rows] == [10000.0, 30000.0]
        frozen = next(row for row in rows if row["cost_sources"][0]["type"] == "FrozenRate")
        assert frozen["material_cost"] == pytest.approx(30_000)
        assert frozen["cost_sources"][0]["qty"] == pytest.approx(30)
        assert frozen["cost_sources"][0]["amount"] == pytest.approx(30_000)
        assert "Frozen sale-line cost" in frozen["cost_basis"]


# --------------------------------------------------------------------------- #
# 9. Booking delivery supplied from two lots
# --------------------------------------------------------------------------- #
def test_booking_delivery_supplied_from_two_lots(multi_app):
    client = make_csrf_client(multi_app)
    login(client)
    with multi_app.app_context():
        material_id, _client_row, supplier_id = seed(multi_app)
        add_grn(multi_app, supplier_id, 10, 1000, datetime(2026, 1, 1, 9, 0), "GRN-OLD")
        add_grn(multi_app, supplier_id, 30, 1200, datetime(2026, 1, 5, 9, 0), "GRN-NEW")

        booking_response = client.post("/add_booking", data={
            "client_name": "CUSTOMER ONE",
            "category": "General",
            "material_name[]": ["CEMENT"],
            "material_id[]": [str(material_id)],
            "qty[]": ["40"],
            "unit_rate[]": ["1300"],
            "date": "2026-01-03",
            "paid_amount": "0",
        }, follow_redirects=False)
        assert booking_response.status_code in (302, 303)

        for qty, date, bill in ((10, "2026-01-06", "BK-DEL-1"), (30, "2026-01-07", "BK-DEL-2")):
            response = client.post("/add_direct_sale", data={
                "client_name": "CUSTOMER ONE",
                "driver_name": "Driver",
                "category": "Booking Delivery",
                "product_name[]": ["CEMENT"],
                "qty[]": [str(qty)],
                "unit_rate[]": ["0"],
                "payment_method": "Credit",
                "paid_amount": "0",
                "manual_bill_no": bill,
                "sale_date": date,
                "allow_negative_stock": "1",
            }, follow_redirects=False)
            assert response.status_code in (302, 303), response.get_data(as_text=True)[:300]

        rep = report()
        supplied = goods_rows(rep, ("booking_supply",))
        assert [(row["qty"], row["material_cost"]) for row in supplied] == [(10.0, 10000.0), (30.0, 36000.0)]
        assert rep.summary["total_cogs"] == pytest.approx(EXPECTED_COGS)
        assert rep.summary["total_material_revenue"] == pytest.approx(52_000)
        assert rep.summary["total_profit"] == pytest.approx(52_000 - EXPECTED_COGS)
        assert {row["cost_sources"][0]["reference"] for row in supplied} == {"MB NO.GRN-OLD", "MB NO.GRN-NEW"}
